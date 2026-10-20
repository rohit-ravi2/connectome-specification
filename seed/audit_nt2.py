"""AUDIT part 2 — the likely killers: (C) contact-stratified NT permutation (correct null),
(D) label specificity: is NT special or just a coarse class label? does NT add beyond
canonical neuron class? All via LR for speed; lifts are relative and internally consistent."""
import openpyxl, csv, re, math, numpy as np
from collections import defaultdict, Counter
from sklearn.model_selection import StratifiedKFold, cross_val_predict
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import roc_auc_score
from scipy import sparse
RNG=np.random.default_rng(1)
DIR="../data/celegans/data/connectome"
T5="../data/tier5/connectome.npz"
# ---- setup (same universe) ----
wb=openpyxl.load_workbook(f"{DIR}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True); ws=wb['adult nerve ring neighbors']
data=list(ws.iter_rows(values_only=True)); neurons=[h for h in data[0][1:]]; idx={n:i for i,n in enumerate(neurons) if isinstance(n,str)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    rn=row[0]
    if rn not in idx: continue
    for j,val in enumerate(row[1:]):
        if val in (None,''): continue
        try: A[idx[rn],j]=float(val)
        except: pass
A=np.maximum(A,A.T); deg_area=A.sum(1); deg_num=(A>0).sum(1)
present=defaultdict(set)
with open(f"{DIR}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
        if ds in (7,8): present[(min(r['pre'],r['post']),max(r['pre'],r['post']))].add(ds)
t5=np.load(T5,allow_pickle=True); t5nt={str(n):str(v) for n,v in zip(t5["names"],t5["nt_primary"])}
print("T5 npz keys:", list(t5.keys()))
def bc(n):
    b=str(n)
    for s in("L","R"):
        if b.endswith(s): b=b[:-1]
    return b
nt_of=[t5nt.get(neurons[i], t5nt.get(bc(neurons[i]),"Unknown")) if isinstance(neurons[i],str) else "Unknown" for i in range(M)]
# canonical class (SI6 longest-prefix)
wbc=openpyxl.load_workbook(f"{DIR}/cook2019/SI 6 Cell class lists.xlsx",read_only=True,data_only=True); wsc=wbc['cell class lists']
classes=set()
for row in wsc.iter_rows(values_only=True):
    for c in row:
        if isinstance(c,str) and re.fullmatch(r'[A-Z]{2,4}[0-9]?',c.strip()): classes.add(c.strip())
classes=sorted(classes,key=len,reverse=True)
def to_class(n):
    if not isinstance(n,str): return "NA"
    for c in classes:
        if n.startswith(c): return c
    return n
cls_of=[to_class(neurons[i]) for i in range(M)]

I,J,y,geom=[],[],[],[]
for i in range(M):
    for j in range(i+1,M):
        if A[i,j]<=0 or not isinstance(neurons[i],str) or not isinstance(neurons[j],str): continue
        s=present.get((min(neurons[i],neurons[j]),max(neurons[i],neurons[j])),set())
        I.append(i);J.append(j); y.append(1 if s=={7,8} else 0)
        vi,vj=A[i],A[j]; cos=float(vi@vj/(np.linalg.norm(vi)*np.linalg.norm(vj)+1e-9)); shared=int(((A[i]>0)&(A[j]>0)).sum())
        geom.append([math.log1p(A[i,j]),cos,shared,math.log1p(min(deg_area[i],deg_area[j])),math.log1p(max(deg_area[i],deg_area[j])),min(deg_num[i],deg_num[j]),max(deg_num[i],deg_num[j])])
I=np.array(I);J=np.array(J);y=np.array(y,float);G=np.array(geom); Gs=(G-G.mean(0))/(G.std(0)+1e-9); N=len(y)
folds=list(StratifiedKFold(5,shuffle=True,random_state=0).split(np.zeros(N),y))
def pairfeat(labels):
    cats=sorted(set(labels)); cix={c:k for k,c in enumerate(cats)}; K=len(cats)
    P=sparse.lil_matrix((N,K*K))
    for k in range(N):
        a,b=cix[labels[I[k]]],cix[labels[J[k]]]; lo,hi=min(a,b),max(a,b); P[k,lo*K+hi]=1
    return P.tocsr()
Gsp=sparse.csr_matrix(Gs)
def auc_of(X):
    p=cross_val_predict(LogisticRegression(max_iter=2000),X,y,cv=folds,method='predict_proba')[:,1]
    return roc_auc_score(y,p)
a_g=auc_of(Gsp); NT=pairfeat(nt_of); CL=pairfeat(cls_of)
a_gn=auc_of(sparse.hstack([Gsp,NT]).tocsr())
a_gc=auc_of(sparse.hstack([Gsp,CL]).tocsr())
a_gnc=auc_of(sparse.hstack([Gsp,NT,CL]).tocsr())
print(f"\nN={N}, reproducible {int(y.sum())} ({100*y.mean():.1f}%)")
print("="*90); print("PART D — LABEL SPECIFICITY: is NT special, or subsumed by canonical class?"); print("="*90)
print(f"  geometry [LR]                 AUC={a_g:.4f}")
print(f"  geometry + NT-pair            AUC={a_gn:.4f}   (lift {a_gn-a_g:+.4f})")
print(f"  geometry + CLASS-pair         AUC={a_gc:.4f}   (lift {a_gc-a_g:+.4f})  [{len(set(cls_of))} classes]")
print(f"  geometry + NT + CLASS         AUC={a_gnc:.4f}")
print(f"  -> NT adds beyond class:      {a_gnc-a_gc:+.4f}   (if ~0, NT is subsumed by class)")
print(f"  -> class adds beyond NT:      {a_gnc-a_gn:+.4f}")

print("\n"+"="*90); print("PART C — CONTACT-STRATIFIED NT PERMUTATION (correct null: shuffle NT among geometrically similar neurons)"); print("="*90)
# strata = contact-degree quartile x contact-area quartile
dq=np.digitize(deg_num,np.quantile(deg_num[[isinstance(neurons[i],str) for i in range(M)]] if False else deg_num,[.25,.5,.75]))
aq=np.digitize(deg_area,np.quantile(deg_area,[.25,.5,.75]))
strata=defaultdict(list)
for i in range(M):
    if isinstance(neurons[i],str): strata[(dq[i],aq[i])].append(i)
real_lift=a_gn-a_g
glob=[]; strat=[]
for _ in range(30):
    # global shuffle
    pl=list(np.array(nt_of)[RNG.permutation(M)])
    glob.append(auc_of(sparse.hstack([Gsp,pairfeat(pl)]).tocsr())-a_g)
    # contact-stratified shuffle
    sl=list(nt_of)
    for st,mem in strata.items():
        perm=RNG.permutation(mem)
        for a,b in zip(mem,perm): sl[a]=nt_of[b]
    strat.append(auc_of(sparse.hstack([Gsp,pairfeat(sl)]).tocsr())-a_g)
glob=np.array(glob); strat=np.array(strat)
print(f"  REAL NT lift over geometry:            {real_lift:+.4f}")
print(f"  global-shuffle null:      mean {glob.mean():+.4f} ± {glob.std():.4f}   z={(real_lift-glob.mean())/(glob.std()+1e-9):.1f}")
print(f"  contact-stratified null:  mean {strat.mean():+.4f} ± {strat.std():.4f}   z={(real_lift-strat.mean())/(strat.std()+1e-9):.1f}")
print(f"  -> if z(stratified) collapses toward 0, the lift is NT-as-contact/position proxy (CONFOUNDED)")
print("="*90)
