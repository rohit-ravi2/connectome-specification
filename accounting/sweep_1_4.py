import openpyxl, csv, re, math, numpy as np
from collections import defaultdict
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.model_selection import cross_val_predict, StratifiedKFold
from sklearn.metrics import roc_auc_score
D="../data/celegans/data/connectome"
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True)); neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); cdeg=(A>0).sum(1)
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1)); dchem[(r['pre'],r['post'])].add(ds)
# Cook directed chemical boolean matrix
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r]))
colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
Cin=C.T@C; Cout=C@C.T; Cff=C@C   # shared-input, shared-output, feedforward (Cook)

# directed contacting pairs
I,J,y7,y8,kL1,feat=[],[],[],[],[],[]
for i in range(M):
    for j in range(M):
        if i==j or A[i,j]<=0: continue
        s=dchem.get((neurons[i],neurons[j]),set())
        I.append(i);J.append(j);y7.append(1 if 7 in s else 0);y8.append(1 if 8 in s else 0)
        kL1.append(sum(1 for d in(1,2,3,4) if d in s))
        feat.append([math.log1p(A[i,j]),int(((A[i]>0)&(A[j]>0)).sum()),float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9)),int(cdeg[i]),int(cdeg[j])])
I=np.array(I);J=np.array(J);y7=np.array(y7);y8=np.array(y8);kL1=np.array(kL1);Xb=np.array(feat);N=len(I)
disagree=(y7!=y8).astype(int); cv=StratifiedKFold(5,shuffle=True,random_state=0)
def h(p): return np.where((p<=0)|(p>=1),0.0,-p*np.log2(np.clip(p,1e-9,1))-(1-p)*np.log2(np.clip(1-p,1e-9,1)))

print("="*74); print("CANDIDATE 1 — LATENT PROPENSITY theta (empirical-Bayes from 4 L1 animals)"); print("="*74)
r=kL1/4.0; m=r.mean(); v=r.var(); prec=m*(1-m)/v-1; al=m*prec; be=(1-m)*prec
theta=(kL1+al)/(4+al+be)
print(f"  Beta prior alpha={al:.2f} beta={be:.2f} | theta distribution:")
print(f"    frac theta<0.1: {(theta<0.1).mean():.2f}  0.1-0.5: {((theta>=0.1)&(theta<0.5)).mean():.2f}  0.5-0.9: {((theta>=0.5)&(theta<0.9)).mean():.2f}  >0.9: {(theta>=0.9).mean():.2f}")
print(f"    -> {'BIMODAL (sharp specification: most pairs near-deterministic)' if (theta<0.1).mean()+(theta>0.9).mean()>0.7 else 'SPREAD (specification is graded)'}")
# does L1-theta predict ADULT (out-of-sample, cross-stage) vs features?
aucF=roc_auc_score(y8,cross_val_predict(GradientBoostingClassifier(random_state=0),Xb,y8,cv=cv,method='predict_proba')[:,1])
aucT=roc_auc_score(y8,theta)
print(f"  predict ADULT-8 edge: features AUC={aucF:.3f}  vs  L1-propensity theta AUC={aucT:.3f}")
# entropy accounting
Hu=h(np.full(N,y8.mean())).mean()
Ht=np.array([h(np.array([t]))[0] for t in np.clip(theta,1e-6,1-1e-6)]).mean()
print(f"  H(adult8) unconditional={Hu:.3f} bits/pair | H given L1-theta={Ht:.3f} | specification removed by theta={Hu-Ht:.3f}")
# prediction: halo (adult disagreement) concentrates at intermediate theta?
print(f"  adult-halo pairs mean theta={theta[disagree==1].mean():.2f} vs core(both) {theta[(y7==1)&(y8==1)].mean():.2f} vs none {theta[(y7==0)&(y8==0)].mean():.2f}")

print("\n"+"="*74); print("CANDIDATE 4 — HIGHER-ORDER wiring context (Cook motif structure, cross-animal)"); print("="*74)
HO=np.column_stack([np.log1p(Cin[I,J]),np.log1p(Cout[I,J]),np.log1p(Cff[I,J]),np.log1p(Cff[J,I])])
Xe=np.hstack([Xb,HO])
def disp(X):
    y2=np.concatenate([y7,y8]); X2=np.vstack([X,X])
    p=np.clip(cross_val_predict(GradientBoostingClassifier(random_state=0),X2,y2,cv=cv,method='predict_proba')[:,1][:N],1e-4,1-1e-4)
    return disagree.mean()/(2*p*(1-p)).mean()
yrep=((y7==1)&(y8==1)).astype(int)
nf={n:i%5 for i,n in enumerate(sorted(set([neurons[i] for i in I]+[neurons[j] for j in J])))}
def nho(X):
    pred=np.full(N,np.nan)
    for f in range(5):
        te=[k for k in range(N) if nf[neurons[I[k]]]==f and nf[neurons[J[k]]]==f]; tr=[k for k in range(N) if nf[neurons[I[k]]]!=f and nf[neurons[J[k]]]!=f]
        if len(te)<20: continue
        pred[np.array(te)]=GradientBoostingClassifier(random_state=0).fit(X[tr],yrep[tr]).predict_proba(X[te])[:,1]
    m=~np.isnan(pred); return roc_auc_score(yrep[m],pred[m])
print(f"  baseline (contact):        dispersion_ratio={disp(Xb):.2f}  AUC(nho)={nho(Xb):.3f}")
print(f"  + Cook higher-order motif: dispersion_ratio={disp(Xe):.2f}  AUC(nho)={nho(Xe):.3f}")
print("  (ratio 0.48 baseline; ->1 = circuit-context captures the sharp propensity)")
