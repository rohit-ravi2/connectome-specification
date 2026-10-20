import openpyxl, csv, re, math, sys, numpy as np
from collections import defaultdict
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.model_selection import cross_val_predict, StratifiedKFold
from sklearn.metrics import roc_auc_score
from scipy.sparse import lil_matrix, csr_matrix
from scipy.sparse.linalg import eigsh
def P(*a): print(*a); sys.stdout.flush()
D="../data/celegans/data/connectome"
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True)); neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); Ab=(A>0).astype(float); cdeg=Ab.sum(1)
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
            if r['pre'] in idx and r['post'] in idx: dchem[(r['pre'],r['post'])].add(ds)
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r])); colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
Cin=C.T@C;Cout=C@C.T;Cff=C@C;Cu=((C+C.T)>0).astype(float);comm=Cu@Cu
Usv,s,Vt=np.linalg.svd(C); dd=16; RD=(Usv[:,:dd]*np.sqrt(s[:dd]))@(np.sqrt(s[:dd])[:,None]*Vt[:dd])
P("data loaded.")
# ---- DEEP: order-3+ multi-hop (Katz on Cook) ----
rho=np.abs(np.linalg.eigvals(C)).max(); beta=0.5/rho
Katz=np.linalg.inv(np.eye(M)-beta*C)-np.eye(M); C3=C@C@C
P(f"Katz built (beta={beta:.3f}).")
# ---- DEEP: multi-scale diffusion geometry of the CONTACT graph ----
deg=Ab.sum(1); Dm=np.diag(1/np.sqrt(np.maximum(deg,1e-9))); Lsym=np.eye(M)-Dm@Ab@Dm
ev,V=np.linalg.eigh(Lsym)  # 182x182, instant
def diffprox(t): W=V*np.exp(-t*ev)[None,:]; return W@V.T   # heat kernel proximity
DP={t:diffprox(t) for t in (1.0,3.0,10.0)}
P("diffusion geometry built.")
# ---- DEEP: simplicial / Hodge structure of the CONTACT complex ----
edges=[(i,j) for i in range(M) for j in range(i+1,M) if Ab[i,j]>0]; eidx={e:k for k,e in enumerate(edges)}; nE=len(edges)
B1=lil_matrix((M,nE))
for (a,b),e in eidx.items(): B1[a,e]=-1; B1[b,e]=1
tris=[]
for a in range(M):
    nb=[x for x in range(M) if Ab[a,x]>0 and x>a]
    for bi in range(len(nb)):
        for ci in range(bi+1,len(nb)):
            b,c=nb[bi],nb[ci]
            if Ab[b,c]>0: tris.append((a,b,c))
B2=lil_matrix((nE,len(tris)))
for t,(a,b,c) in enumerate(tris):
    B2[eidx[(a,b)],t]=1; B2[eidx[(a,c)],t]=-1; B2[eidx[(b,c)],t]=1
B1=B1.tocsr(); B2=B2.tocsr(); L1=(B1.T@B1+B2@B2.T).tocsr()
P(f"Hodge L1 built: {nE} edges, {len(tris)} triangles.")
kH=24
try:
    hev,HV=eigsh(L1,k=kH,which='SA',maxiter=5000); order=np.argsort(hev); HV=HV[:,order]
    P(f"Hodge spectrum: {kH} low eigenvectors (smallest eig {hev.min():.3f}).")
    hodge_ok=True
except Exception as e:
    P("Hodge eigsh failed:",e); hodge_ok=False; HV=np.zeros((nE,kH))
# ---- pair table ----
I,J,y7,y8,base,katz,spec,hodge=[],[],[],[],[],[],[],[]
for i in range(M):
    for j in range(M):
        if i==j or A[i,j]<=0: continue
        st=dchem.get((neurons[i],neurons[j]),set())
        I.append(i);J.append(j);y7.append(1 if 7 in st else 0);y8.append(1 if 8 in st else 0)
        base.append([math.log1p(A[i,j]),int(((A[i]>0)&(A[j]>0)).sum()),float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9)),int(cdeg[i]),int(cdeg[j]),
                     np.log1p(Cin[i,j]),np.log1p(Cout[i,j]),np.log1p(Cff[i,j]),np.log1p(comm[i,j]),RD[i,j],RD[j,i]])
        katz.append([Katz[i,j],Katz[j,i],np.log1p(C3[i,j]),np.log1p(C3[j,i])])
        spec.append([DP[1.0][i,j],DP[3.0][i,j],DP[10.0][i,j]])
        e=eidx[(min(i,j),max(i,j))]; hodge.append(list(HV[e]))
I=np.array(I);J=np.array(J);y7=np.array(y7);y8=np.array(y8);N=len(I)
base=np.array(base);katz=np.array(katz);spec=np.array(spec);hodge=np.array(hodge)
disagree=(y7!=y8).astype(int);yrep=((y7==1)&(y8==1)).astype(int);cv=StratifiedKFold(5,shuffle=True,random_state=0)
nf={n:i%5 for i,n in enumerate(sorted(set([neurons[i] for i in I]+[neurons[j] for j in J])))}
def nho(X,y=yrep):
    pred=np.full(N,np.nan)
    for f in range(5):
        te=[k for k in range(N) if nf[neurons[I[k]]]==f and nf[neurons[J[k]]]==f];tr=[k for k in range(N) if nf[neurons[I[k]]]!=f and nf[neurons[J[k]]]!=f]
        if len(te)<20: continue
        pred[np.array(te)]=GradientBoostingClassifier(random_state=0).fit(X[tr],y[tr]).predict_proba(X[te])[:,1]
    m=~np.isnan(pred);return roc_auc_score(y[m],pred[m])
def disp(X):
    y2=np.concatenate([y7,y8]);X2=np.vstack([X,X]);p=np.clip(cross_val_predict(GradientBoostingClassifier(random_state=0),X2,y2,cv=cv,method='predict_proba')[:,1][:N],1e-4,1-1e-4)
    return disagree.mean()/(2*p*(1-p)).mean()
# ---- NOISE CEILING (reference only) ----
p1=y8[y7==1].mean(); p0=y8[y7==0].mean(); ceil_auc=roc_auc_score(y8,np.where(y7==1,p1,p0))
P(f"\nNOISE CEILING (adult7->adult8, ref line): AUC={ceil_auc:.3f}  | disagreement rate={disagree.mean():.3f}")
P("="*70)
b_disp,b_auc=disp(base),nho(base); P(f"baseline (contact+HO+RDPG16):     disp={b_disp:.2f}  AUC={b_auc:.3f}")
for name,F in [("+Katz/3-hop (order-3+ multi-hop)",katz),("+diffusion geometry (multi-scale)",spec),("+Hodge simplicial (cycles/holes)",hodge)]:
    X=np.hstack([base,F]); P(f"{name:34s} disp={disp(X):.2f}  AUC={nho(X):.3f}  (Δ={nho(X)-b_auc:+.3f})")
Xall=np.hstack([base,katz,spec,hodge]); P(f"{'+ALL deep layers stacked':34s} disp={disp(Xall):.2f}  AUC={nho(Xall):.3f}  (Δ={nho(Xall)-b_auc:+.3f})")
P(f"\n  reference: noise-ceiling AUC {ceil_auc:.3f} — headroom above baseline = {ceil_auc-b_auc:+.3f}")
