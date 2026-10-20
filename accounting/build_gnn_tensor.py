import openpyxl, re, sys, numpy as np
from collections import defaultdict
def P(*a): print(*a); sys.stdout.flush()
D="../data/celegans/data/connectome"
SC="../work"
z=np.load(f"{SC}/deepfeats.npz",allow_pickle=True)
I,J,y7,y8,base,katz,neurons=z['I'],z['J'],z['y7'],z['y8'],z['base'],z['katz'],list(z['neurons']); idx={n:i for i,n in enumerate(neurons)}; M=len(neurons); N=len(I)
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r])); colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
U,s,Vt=np.linalg.svd(C); dd=16; UL=U[:,:dd]*np.sqrt(s[:dd]); VR=(np.sqrt(s[:dd])[:,None]*Vt[:dd]).T
# contact matrix for attributes
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True))
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); Ab=(A>0).astype(float); cdeg=Ab.sum(1)
# NT labels
t5=np.load("../data/tier5/connectome.npz",allow_pickle=True)
t5nt={str(n):str(v) for n,v in zip(t5["names"],t5["nt_primary"])}
def bc(n):
    b=str(n)
    for sfx in("L","R"):
        if b.endswith(sfx): b=b[:-1]
    return b
nts=[t5nt.get(neurons[i],t5nt.get(bc(neurons[i]),"Unk")) for i in range(M)]
cats=sorted(set(nts)); NToh=np.zeros((M,len(cats)))
for i,n in enumerate(nts): NToh[i,cats.index(n)]=1
# ---- (A) nonlinear inductive random-weight GNN embedding on Cook connectome ----
Csym=((C+C.T)>0).astype(float); Asym=Csym+np.eye(M); dsm=Asym.sum(1); An=np.diag(1/np.sqrt(dsm))@Asym@np.diag(1/np.sqrt(dsm))
X0=np.hstack([np.log1p(cdeg)[:,None],np.log1p(C.sum(1))[:,None],np.log1p(C.sum(0))[:,None],UL,VR,NToh])
X0=(X0-X0.mean(0))/(X0.std(0)+1e-9)
rng=np.random.default_rng(0)
W1=rng.normal(0,1/np.sqrt(X0.shape[1]),(X0.shape[1],48)); H1=np.tanh(An@X0@W1)
W2=rng.normal(0,1/np.sqrt(48),(48,48)); H2=np.tanh(An@H1@W2)   # 2-hop nonlinear
Z=np.hstack([H1,H2]); Z=(Z-Z.mean(0))/(Z.std(0)+1e-9)
Uz,sz,_=np.linalg.svd(Z-Z.mean(0),full_matrices=False); Zr=(Uz[:,:16]*sz[:16])   # PCA to 16
# ---- (B) explicit 3-way tensor / triadic features ----
Snl=np.tanh(UL@UL.T); Tnl=np.tanh(UL@VR.T); TRI=Snl@Tnl  # nonlinear trilinear aggregate sum_k s(Ui.Uk)s(Uk.Vj)
gnn=np.zeros((N,18)); tens=np.zeros((N,7))
for r,(i,j) in enumerate(zip(I,J)):
    zi,zj=Zr[i],Zr[j]; gnn[r]=np.concatenate([zi*zj,[zi@zj,np.linalg.norm(zi-zj)]])
    ff=C[i,:]*C[:,j]; cs=C[:,i]*C[:,j]; ct=C[i,:]*C[j,:]
    tens[r]=[ff.max(),np.log1p((ff>0).sum()),ff.max()/(ff.sum()+1e-9),cs.max(),ct.max(),TRI[i,j],TRI[j,i]]
np.savez(f"{SC}/gnntensor.npz",gnn=gnn,tens=tens)
P("saved gnntensor.npz",gnn.shape,tens.shape,"| NT classes",len(cats))
