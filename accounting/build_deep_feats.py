import openpyxl, csv, re, math, sys, numpy as np
from collections import defaultdict
from scipy.sparse import lil_matrix
from scipy.sparse.linalg import eigsh
def P(*a): print(*a); sys.stdout.flush()
D="../data/celegans/data/connectome"
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True)); neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); Ab=(A>0).astype(float); cdeg=Ab.sum(1)
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
            if r['pre'] in idx and r['post'] in idx: dchem[(r['pre'],r['post'])].add(ds)
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r])); colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
Cin=C.T@C;Cout=C@C.T;Cff=C@C;Cu=((C+C.T)>0).astype(float);comm=Cu@Cu
Usv,s,Vt=np.linalg.svd(C); dd=16; RD=(Usv[:,:dd]*np.sqrt(s[:dd]))@(np.sqrt(s[:dd])[:,None]*Vt[:dd])
rho=np.abs(np.linalg.eigvals(C)).max(); beta=0.5/rho
Katz=np.linalg.inv(np.eye(M)-beta*C)-np.eye(M); C3=C@C@C; C4=C3@C
deg=Ab.sum(1); Dm=np.diag(1/np.sqrt(np.maximum(deg,1e-9))); Lsym=np.eye(M)-Dm@Ab@Dm
ev,V=np.linalg.eigh(Lsym)
def diffprox(t): W=V*np.exp(-t*ev)[None,:]; return W@V.T
DP={t:diffprox(t) for t in (1.0,3.0,10.0)}
# Hodge on contact complex + higher kH
edges=[(i,j) for i in range(M) for j in range(i+1,M) if Ab[i,j]>0]; eidx={e:k for k,e in enumerate(edges)}; nE=len(edges)
B1=lil_matrix((M,nE))
for (a,b),e in eidx.items(): B1[a,e]=-1; B1[b,e]=1
tris=[]
for a in range(M):
    nb=[x for x in range(M) if Ab[a,x]>0 and x>a]
    for bi in range(len(nb)):
        for ci in range(bi+1,len(nb)):
            b,c=nb[bi],nb[ci]
            if Ab[b,c]>0: tris.append((a,b,c))
B2=lil_matrix((nE,len(tris)))
for t,(a,b,c) in enumerate(tris):
    B2[eidx[(a,b)],t]=1; B2[eidx[(a,c)],t]=-1; B2[eidx[(b,c)],t]=1
B1=B1.tocsr(); B2=B2.tocsr(); L1=(B1.T@B1+B2@B2.T).tocsr()
kH=40; hev,HV=eigsh(L1,k=kH,which='SA',maxiter=8000); o=np.argsort(hev); HV=HV[:,o]; hev=hev[o]
P(f"Hodge: {nE} edges {len(tris)} tris; smallest eigs {hev[:4].round(3)}")
# also 4-hop Katz-like higher order already in Katz; add motif clustering (triangle participation per node)
tri_deg=np.zeros(M)
for a,b,c in tris: tri_deg[a]+=1; tri_deg[b]+=1; tri_deg[c]+=1
I,J,y7,y8,base,katz,spec,hodge=[],[],[],[],[],[],[],[]
for i in range(M):
    for j in range(M):
        if i==j or A[i,j]<=0: continue
        st=dchem.get((neurons[i],neurons[j]),set())
        I.append(i);J.append(j);y7.append(1 if 7 in st else 0);y8.append(1 if 8 in st else 0)
        base.append([math.log1p(A[i,j]),int(((A[i]>0)&(A[j]>0)).sum()),float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9)),int(cdeg[i]),int(cdeg[j]),
                     np.log1p(Cin[i,j]),np.log1p(Cout[i,j]),np.log1p(Cff[i,j]),np.log1p(comm[i,j]),RD[i,j],RD[j,i]])
        katz.append([Katz[i,j],Katz[j,i],np.log1p(C3[i,j]),np.log1p(C3[j,i]),np.log1p(C4[i,j]),np.log1p(C4[j,i])])
        spec.append([DP[1.0][i,j],DP[3.0][i,j],DP[10.0][i,j],tri_deg[i],tri_deg[j]])
        e=eidx[(min(i,j),max(i,j))]; hodge.append(list(HV[e]))
np.savez("../work/deepfeats.npz",
         I=np.array(I),J=np.array(J),y7=np.array(y7),y8=np.array(y8),
         base=np.array(base),katz=np.array(katz),spec=np.array(spec),hodge=np.array(hodge),
         neurons=np.array(neurons,dtype=object))
P("saved deepfeats.npz  shapes:",np.array(base).shape,np.array(katz).shape,np.array(hodge).shape)
