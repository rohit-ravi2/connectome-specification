import openpyxl, csv, re, math, numpy as np
from collections import defaultdict
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.model_selection import cross_val_predict, StratifiedKFold
from sklearn.metrics import roc_auc_score
D="../data/celegans/data/connectome"
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True)); neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); cdeg=(A>0).sum(1)
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
            if r['pre'] in idx and r['post'] in idx: dchem[(r['pre'],r['post'])].add(ds)
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r])); colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
# RDPG of Cook
U,s,Vt=np.linalg.svd(C)
# L1 developmental graph (union ds1-4)
L1g=np.zeros((M,M))
for (a,b),st in dchem.items():
    if st&{1,2,3,4}: L1g[idx[a],idx[b]]=1
L1in=L1g.T@L1g; L1out=L1g@L1g.T; L1ff=L1g@L1g
# Cook higher-order (winner baseline)
Cin=C.T@C;Cout=C@C.T;Cff=C@C;Cu=((C+C.T)>0).astype(float);comm=Cu@Cu
I,J,y7,y8,feat=[],[],[],[],[]
for i in range(M):
    for j in range(M):
        if i==j or A[i,j]<=0: continue
        st=dchem.get((neurons[i],neurons[j]),set())
        I.append(i);J.append(j);y7.append(1 if 7 in st else 0);y8.append(1 if 8 in st else 0)
        feat.append([math.log1p(A[i,j]),int(((A[i]>0)&(A[j]>0)).sum()),float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9)),int(cdeg[i]),int(cdeg[j])])
I=np.array(I);J=np.array(J);y7=np.array(y7);y8=np.array(y8);Xb=np.array(feat);N=len(I)
disagree=(y7!=y8).astype(int);yrep=((y7==1)&(y8==1)).astype(int);cv=StratifiedKFold(5,shuffle=True,random_state=0)
HO=np.column_stack([np.log1p(Cin[I,J]),np.log1p(Cout[I,J]),np.log1p(Cff[I,J]),np.log1p(comm[I,J])])
nf={n:i%5 for i,n in enumerate(sorted(set([neurons[i] for i in I]+[neurons[j] for j in J])))}
def nho(X,y=yrep):
    pred=np.full(N,np.nan)
    for f in range(5):
        te=[k for k in range(N) if nf[neurons[I[k]]]==f and nf[neurons[J[k]]]==f];tr=[k for k in range(N) if nf[neurons[I[k]]]!=f and nf[neurons[J[k]]]!=f]
        if len(te)<20: continue
        pred[np.array(te)]=GradientBoostingClassifier(random_state=0).fit(X[tr],y[tr]).predict_proba(X[te])[:,1]
    m=~np.isnan(pred);return roc_auc_score(y[m],pred[m])
def disp(X):
    y2=np.concatenate([y7,y8]);X2=np.vstack([X,X]);p=np.clip(cross_val_predict(GradientBoostingClassifier(random_state=0),X2,y2,cv=cv,method='predict_proba')[:,1][:N],1e-4,1-1e-4)
    return disagree.mean()/(2*p*(1-p)).mean()
print(f"baseline contact: disp={disp(Xb):.2f} AUC={nho(Xb):.3f} | contact+HO: disp={disp(np.hstack([Xb,HO])):.2f} AUC={nho(np.hstack([Xb,HO])):.3f}")
print("\n#2 LATENT-SPACE (Cook RDPG rank sweep) — is the propensity low-rank latent?")
for d in [2,4,8,16,32]:
    Ld=(U[:,:d]*np.sqrt(s[:d]))@ (np.sqrt(s[:d])[:,None]*Vt[:d])  # rank-d reconstruction of C
    LAT=np.column_stack([Ld[I,J],Ld[J,I]])
    print(f"  rank d={d:2d}: contact+RDPG  disp={disp(np.hstack([Xb,LAT])):.2f} AUC={nho(np.hstack([Xb,LAT])):.3f}")
print("\n#1 GROWTH PROCESS — does the EARLY (L1) graph structure predict adult edges?")
GR=np.column_stack([np.log1p(L1in[I,J]),np.log1p(L1out[I,J]),np.log1p(L1ff[I,J])])
print(f"  contact + L1-graph context: disp={disp(np.hstack([Xb,GR])):.2f} AUC={nho(np.hstack([Xb,GR])):.3f}")
print("\n#3 SELF-AVOIDANCE — among STRONG contacts, is non-wiring reproducibly specified (repulsion)?")
area=Xb[:,0]; strong=area>np.quantile(area,0.75)
ys,y7s,y8s=yrep[strong],y7[strong],y8[strong]
# among strong contacts: edge rate, and reproducibility of NON-wiring
sc=np.where(strong)[0]
# dispersion of wiring within strong contacts
def dispsub(mask):
    a7,a8=y7[mask],y8[mask]; d=(a7!=a8).mean()
    p=np.clip(cross_val_predict(GradientBoostingClassifier(random_state=0),np.vstack([Xb[mask],Xb[mask]]),np.concatenate([a7,a8]),cv=StratifiedKFold(5,shuffle=True,random_state=0),method='predict_proba')[:,1][:mask.sum()],1e-4,1-1e-4)
    return d/(2*p*(1-p)).mean()
nonwire_both=((y7[strong]==0)&(y8[strong]==0)).mean(); wire_both=((y7[strong]==1)&(y8[strong]==1)).mean(); variable=((y7[strong]!=y8[strong])).mean()
print(f"  strong-contact pairs: {strong.sum()} | wired-both={wire_both:.2f} NON-wired-both(repulsion)={nonwire_both:.2f} variable={variable:.2f}")
print(f"  -> {nonwire_both/(nonwire_both+variable+1e-9)*100:.0f}% of non-wired strong contacts are REPRODUCIBLY non-wired (specified repulsion vs stochastic)")
print(f"  predict wiring among strong contacts: contact-only AUC={nho(Xb[strong] if False else Xb, y=yrep):.3f} (full); +HO on strong subset:")
aHO_s=roc_auc_score(ys,cross_val_predict(GradientBoostingClassifier(random_state=0),np.hstack([Xb,HO])[strong],ys,cv=cv,method='predict_proba')[:,1])
ac_s=roc_auc_score(ys,cross_val_predict(GradientBoostingClassifier(random_state=0),Xb[strong],ys,cv=cv,method='predict_proba')[:,1])
print(f"    within strong contacts: contact AUC={ac_s:.3f} -> +HO {aHO_s:.3f} (can we predict which strong contacts wire vs are repelled?)")
