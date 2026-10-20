"""Is the class-pair code just contact geometry in disguise, or complementary?
Compare geometry / class-pair / combined on IDENTICAL CV folds."""
import openpyxl, csv, re, math, numpy as np
from collections import defaultdict
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.model_selection import StratifiedKFold
from sklearn.metrics import roc_auc_score

DIR="../data/celegans/data/connectome"
wb=openpyxl.load_workbook(f"{DIR}/cook2019/SI 6 Cell class lists.xlsx",read_only=True,data_only=True); ws=wb['cell class lists']
classes=set()
for row in ws.iter_rows(values_only=True):
    for c in row:
        if isinstance(c,str) and re.fullmatch(r'[A-Z]{2,4}[0-9]?',c.strip()): classes.add(c.strip())
classes=sorted(classes,key=len,reverse=True)
def to_class(n):
    if not isinstance(n,str): return "NA"
    for c in classes:
        if n.startswith(c): return c
    return n
wb2=openpyxl.load_workbook(f"{DIR}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
ws2=wb2['adult nerve ring neighbors']; data=list(ws2.iter_rows(values_only=True))
neurons=[h for h in data[0][1:]]; idx={n:i for i,n in enumerate(neurons) if isinstance(n,str)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    rn=row[0]
    if rn not in idx: continue
    for j,val in enumerate(row[1:]):
        if val in (None,''): continue
        try: A[idx[rn],j]=float(val)
        except (TypeError,ValueError): pass
A=np.maximum(A,A.T); deg_area=A.sum(1); deg_num=(A>0).sum(1); ncls=[to_class(n) for n in neurons]
chem=defaultdict(set); elec=defaultdict(set)
with open(f"{DIR}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
        if ds not in (7,8): continue
        pre,post,typ=r['pre'],r['post'],r['type']
        (elec if typ=='electrical' else chem)[((min(pre,post),max(pre,post)) if typ=='electrical' else (pre,post))].add(ds)
repro=set()
for e,dss in chem.items():
    if len(dss)==2: repro.add((min(e),max(e)))
for e,dss in elec.items():
    if len(dss)==2: repro.add(e)
Xg=[]; cps=[]; y=[]
for i in range(M):
    for j in range(i+1,M):
        if A[i,j]<=0 or not isinstance(neurons[i],str) or not isinstance(neurons[j],str): continue
        vi,vj=A[i],A[j]; cos=float(vi@vj/(np.linalg.norm(vi)*np.linalg.norm(vj)+1e-9))
        Xg.append([math.log1p(A[i,j]),cos,int(((A[i]>0)&(A[j]>0)).sum()),
                   math.log1p(min(deg_area[i],deg_area[j])),math.log1p(max(deg_area[i],deg_area[j])),
                   min(deg_num[i],deg_num[j]),max(deg_num[i],deg_num[j])])
        cps.append(tuple(sorted((ncls[i],ncls[j])))); y.append(1 if (neurons[i],neurons[j]) in repro or (neurons[j],neurons[i]) in repro else 0)
Xg=np.array(Xg); y=np.array(y,float); N=len(y); pbar=y.mean()
cv=StratifiedKFold(5,shuffle=True,random_state=0)
def evalp(p): return roc_auc_score(y,p), -(y*np.log2(np.clip(p,1e-6,1-1e-6))+(1-y)*np.log2(np.clip(1-p,1e-6,1-1e-6))).sum()
pg=np.zeros(N); pc=np.zeros(N); pgc=np.zeros(N); alpha=5.0
for tr,te in cv.split(Xg,y):
    # class target-encoding (smoothed), fit on train
    tab=defaultdict(lambda:[0.0,0.0])
    for k in tr: tab[cps[k]][0]+=1; tab[cps[k]][1]+=y[k]
    enc=np.array([ (lambda ne:(ne[1]+alpha*pbar)/(ne[0]+alpha))(tab.get(cps[k],[0.0,0.0])) for k in range(N)])
    pc[te]=enc[te]
    GradientBoostingClassifier(random_state=0).fit(Xg[tr],y[tr]); # geometry
    g=GradientBoostingClassifier(random_state=0).fit(Xg[tr],y[tr]); pg[te]=g.predict_proba(Xg[te])[:,1]
    Xc=np.column_stack([Xg,enc])   # geometry + class-encoding
    gc=GradientBoostingClassifier(random_state=0).fit(Xc[tr],y[tr]); pgc[te]=gc.predict_proba(Xc[te])[:,1]
for name,p in [("geometry only",pg),("class-pair only",pc),("geometry + class-pair",pgc)]:
    a,r=evalp(p); print(f"  {name:24s}: AUC {a:.3f}   residual {r:.0f} bits ({r/M:.2f}/neuron)")
