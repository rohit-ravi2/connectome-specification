import openpyxl, csv, re, math, numpy as np
from collections import defaultdict
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.model_selection import cross_val_predict, StratifiedKFold
D="../data/celegans/data/connectome"
# contact
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True)); neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); cdeg=(A>0).sum(1)
# Witvliet chemical directed
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1)); dchem[(r['pre'],r['post'])].add(ds)
# Cook chemical directed (SI5 hermaphrodite chemical)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
# find header row = row with most neuron-name cells
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
hr=max(range(min(6,len(cd))),key=lambda r:sum(isneu(x) for x in cd[r]))
colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
# row header col = the column index that holds neuron names down the rows
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
cook=defaultdict(float)
for r in range(hr+1,len(cd)):
    if rhc>=len(cd[r]) or not isneu(cd[r][rhc]): continue
    pre=cd[r][rhc].strip()
    for c,post in colmap.items():
        if c<len(cd[r]) and cd[r][c] not in (None,''):
            try: cook[(pre,post)]=float(cd[r][c])
            except: pass
cookset=set([k[0] for k in cook]+[k[1] for k in cook])
print(f"Cook chemical directed entries: {len(cook)}; Cook nerve-ring neurons matched: {len(cookset & set(neurons))}")

# directed contacting pairs where both neurons in Cook set
I,J,feat,a7,a8,ck=[],[],[],[],[],[]
for i in range(M):
    for j in range(M):
        if i==j or A[i,j]<=0: continue
        ni,nj=neurons[i],neurons[j]
        if ni not in cookset or nj not in cookset: continue
        s=dchem.get((ni,nj),set())
        feat.append([math.log1p(A[i,j]),int(((A[i]>0)&(A[j]>0)).sum()),float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9)),int(cdeg[i]),int(cdeg[j])])
        a7.append(1 if 7 in s else 0); a8.append(1 if 8 in s else 0); ck.append(1 if cook.get((ni,nj),0)>=1 else 0)
X=np.array(feat); a7=np.array(a7); a8=np.array(a8); ck=np.array(ck); N=len(X)
print(f"contacting directed pairs (both in Cook): {N}  | edge rates W7={a7.mean():.3f} W8={a8.mean():.3f} Cook={ck.mean():.3f}")
cv=StratifiedKFold(5,shuffle=True,random_state=0)
def ratio(recons):
    yp=np.concatenate(recons); Xp=np.vstack([X]*len(recons))
    p=np.clip(cross_val_predict(GradientBoostingClassifier(random_state=0),Xp,yp,cv=cv,method='predict_proba')[:,1][:N],1e-4,1-1e-4)
    k=np.sum(recons,axis=0); n=len(recons); npair=n*(n-1)/2
    return (k*(n-k)/npair).mean(), (2*p*(1-p)).mean()
print("\nDISPERSION (ratio<1 = under-dispersed = agree more than feature-sampling predicts)")
for lab,rc in [("Witvliet W7 vs W8 (same pipeline)",[a7,a8]),
               ("Witvliet-adult vs COOK (CROSS-pipeline)",[a8,ck]),
               ("all three W7,W8,Cook",[a7,a8,ck])]:
    o,pr=ratio(rc); print(f"  {lab:40s} obs={o:.3f} pred={pr:.3f} ratio={o/pr:.2f}")
print("\n  caveat: Cook is different individual+EM era; if it MISSES weak edges its disagreement inflates")
print("  -> if CROSS-pipeline ratio still <1, the under-dispersion is real biology, not Witvliet-pipeline artifact")
