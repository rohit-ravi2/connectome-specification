"""Part 0: master edge table for specification accounting. All nerve-ring pairs
x {contact, developmental trajectory, topology, class/NT, targets}."""
import openpyxl, csv, re, numpy as np
from collections import defaultdict
D="../data/celegans/data/connectome"; T5="../data/tier5/connectome.npz"
# --- contact matrix ---
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
ws=wb['adult nerve ring neighbors']; data=list(ws.iter_rows(values_only=True))
neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    rn=row[0]
    if rn not in idx: continue
    for j,val in enumerate(row[1:]):
        if val not in (None,''):
            try: A[idx[rn],j]=float(val)
            except: pass
A=np.maximum(A,A.T); cdeg=(A>0).sum(1)
# --- class + NT ---
wbc=openpyxl.load_workbook(f"{D}/cook2019/SI 6 Cell class lists.xlsx",read_only=True,data_only=True)
classes=set()
for r in wbc['cell class lists'].iter_rows(values_only=True):
    for c in r:
        if isinstance(c,str) and re.fullmatch(r'[A-Z]{2,4}[0-9]?',c.strip()): classes.add(c.strip())
classes=sorted(classes,key=len,reverse=True)
def cls(n):
    for c in classes:
        if n.startswith(c): return c
    return n
t5=np.load(T5,allow_pickle=True); t5nt={str(n):str(v) for n,v in zip(t5["names"],t5["nt_primary"])}
def base(n):
    b=n
    for s in("L","R"):
        if b.endswith(s): b=b[:-1]
    return b
NT=[t5nt.get(neurons[i],t5nt.get(base(neurons[i]),"Unknown")) for i in range(M)]
CL=[cls(n) for n in neurons]
# --- Witvliet per-stage presence + synapses ---
stage={1:'L1',2:'L1',3:'L1',4:'L1',5:'L2',6:'L3',7:'ad7',8:'ad8'}
pres=defaultdict(set); syn=defaultdict(lambda: defaultdict(int))
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
        pr=(min(r['pre'],r['post']),max(r['pre'],r['post'])); pres[pr].add(ds); syn[pr][ds]+=int(r['synapses'])
# --- build all pairs ---
rows=[]
for i in range(M):
    for j in range(i+1,M):
        ni,nj=neurons[i],neurons[j]; pr=(min(ni,nj),max(ni,nj)); s=pres.get(pr,set())
        vi,vj=A[i]>0,A[j]>0; shared=int((vi&vj).sum())
        cos=float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9))
        a7=1 if 7 in s else 0; a8=1 if 8 in s else 0
        L1=sum(1 for d in(1,2,3,4) if d in s)  # 0-4 of the L1 animals
        rows.append(dict(i=i,j=j,src=ni,tgt=nj,src_class=CL[i],tgt_class=CL[j],src_nt=NT[i],tgt_nt=NT[j],
            contact=int(A[i,j]>0),area=A[i,j],log_area=np.log1p(A[i,j]),shared_nb=shared,cos_sim=round(cos,4),
            cdeg_min=int(min(cdeg[i],cdeg[j])),cdeg_max=int(max(cdeg[i],cdeg[j])),
            same_class=int(CL[i]==CL[j]),lr_homolog=int(base(ni)==base(nj) and ni!=nj),
            L1n=L1,L2=int(5 in s),L3=int(6 in s),ad7=a7,ad8=a8,
            reproducible=int(a7 and a8),halo=int(a7^a8),
            syn7=syn[pr].get(7,0),syn8=syn[pr].get(8,0),synmin=min(syn[pr].get(7,0),syn[pr].get(8,0))))
import csv as _c
keys=list(rows[0].keys())
with open("data/processed/specification_accounting/master_edge_table.csv","w",newline="") as f:
    w=_c.DictWriter(f,keys); w.writeheader(); w.writerows(rows)
N=len(rows); nc=sum(r['contact'] for r in rows); rep=sum(r['reproducible'] for r in rows); halo=sum(r['halo'] for r in rows)
core=sum(1 for r in rows if r['reproducible'] and r['synmin']>=2)
print(f"pairs={N} neurons={M} contacting={nc} reproducible={rep} weak-halo={halo} strong-core(repro&synmin>=2)={core}")
# --- NOISE FLOOR: H(ad8|ad7) over contacting pairs (replicate conditional entropy) ---
import math
def h(p): return 0.0 if p<=0 or p>=1 else -p*math.log2(p)-(1-p)*math.log2(1-p)
C=[r for r in rows if r['contact']]
a7=np.array([r['ad7'] for r in C]); a8=np.array([r['ad8'] for r in C])
# H(ad8|ad7)
floor=0.0
for v in (0,1):
    m=a7==v
    if m.sum(): floor+=m.mean()*h(a8[m].mean())
base_ent=h(a8.mean())
print(f"NOISE FLOOR H(ad8|ad7)={floor:.4f} bits/pair  (base H(ad8)={base_ent:.4f})  reducible={base_ent-floor:.4f}")
print(f"replicate agreement: P(ad8|ad7=1)={a8[a7==1].mean():.3f} P(ad8|ad7=0)={a8[a7==0].mean():.3f}")
