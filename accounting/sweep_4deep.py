import openpyxl, csv, re, math, numpy as np
from collections import defaultdict
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.model_selection import cross_val_predict, StratifiedKFold
from sklearn.metrics import roc_auc_score
D="../data/celegans/data/connectome"
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True)); neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); cdeg=(A>0).sum(1)
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1)); dchem[(r['pre'],r['post'])].add(ds)
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r])); colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
Cu=((C+C.T)>0).astype(float)  # Cook undirected
Cin=C.T@C; Cout=C@C.T; Cff=C@C; comm=Cu@Cu  # common neighbors (undirected)
# Witvliet adult-7 graph for internal cross-check
W7=np.zeros((M,M))
for (a,b),s in dchem.items():
    if 7 in s and a in idx and b in idx: W7[idx[a],idx[b]]=1
W7u=((W7+W7.T)>0).astype(float); W7comm=W7u@W7u; W7in=W7.T@W7; W7out=W7@W7.T; W7ff=W7@W7
I,J,y7,y8,kL1,feat=[],[],[],[],[],[]
for i in range(M):
    for j in range(M):
        if i==j or A[i,j]<=0: continue
        s=dchem.get((neurons[i],neurons[j]),set())
        I.append(i);J.append(j);y7.append(1 if 7 in s else 0);y8.append(1 if 8 in s else 0);kL1.append(sum(1 for d in(1,2,3,4) if d in s))
        feat.append([math.log1p(A[i,j]),int(((A[i]>0)&(A[j]>0)).sum()),float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9)),int(cdeg[i]),int(cdeg[j])])
I=np.array(I);J=np.array(J);y7=np.array(y7);y8=np.array(y8);kL1=np.array(kL1);Xb=np.array(feat);N=len(I)
disagree=(y7!=y8).astype(int); cv=StratifiedKFold(5,shuffle=True,random_state=0)
r=kL1/4.0; m=r.mean(); v=r.var(); prec=m*(1-m)/v-1; theta=(kL1+m*prec)/(4+prec)
HO=np.column_stack([np.log1p(Cin[I,J]),np.log1p(Cout[I,J]),np.log1p(Cff[I,J]),np.log1p(Cff[J,I]),np.log1p(comm[I,J])])
HOw7=np.column_stack([np.log1p(W7in[I,J]),np.log1p(W7out[I,J]),np.log1p(W7ff[I,J]),np.log1p(W7comm[I,J])])
yrep=((y7==1)&(y8==1)).astype(int)
nf={n:i%5 for i,n in enumerate(sorted(set([neurons[i] for i in I]+[neurons[j] for j in J])))}
def nho(X,y=yrep):
    pred=np.full(N,np.nan)
    for f in range(5):
        te=[k for k in range(N) if nf[neurons[I[k]]]==f and nf[neurons[J[k]]]==f]; tr=[k for k in range(N) if nf[neurons[I[k]]]!=f and nf[neurons[J[k]]]!=f]
        if len(te)<20: continue
        pred[np.array(te)]=GradientBoostingClassifier(random_state=0).fit(X[tr],y[tr]).predict_proba(X[te])[:,1]
    mm=~np.isnan(pred); return roc_auc_score(y[mm],pred[mm])
def disp(X):
    y2=np.concatenate([y7,y8]); X2=np.vstack([X,X])
    p=np.clip(cross_val_predict(GradientBoostingClassifier(random_state=0),X2,y2,cv=cv,method='predict_proba')[:,1][:N],1e-4,1-1e-4)
    return disagree.mean()/(2*p*(1-p)).mean()
print("DEEPENING CANDIDATE 4 (higher-order) + combine with latent theta")
print(f"  baseline contact:                 disp={disp(Xb):.2f} AUC={nho(Xb):.3f}")
print(f"  + Cook HO (richer, 5 feats):      disp={disp(np.hstack([Xb,HO])):.2f} AUC={nho(np.hstack([Xb,HO])):.3f}")
print(f"  + Witvliet-INTERNAL HO (W7->W8):  disp={disp(np.hstack([Xb,HOw7])):.2f} AUC={nho(np.hstack([Xb,HOw7])):.3f}")
print(f"  + Cook HO + W7 HO (all context):  disp={disp(np.hstack([Xb,HO,HOw7])):.2f} AUC={nho(np.hstack([Xb,HO,HOw7])):.3f}")
# degree-control: permute HO within degree-quintile bins, does gain survive?
dq=np.digitize(Xb[:,3]+Xb[:,4],np.quantile(Xb[:,3]+Xb[:,4],[.2,.4,.6,.8]))
rng=np.random.default_rng(0); HOperm=HO.copy()
for b in np.unique(dq):
    ix=np.where(dq==b)[0]; HOperm[ix]=HO[rng.permutation(ix)]
print(f"  + Cook HO degree-permuted (null):  AUC={nho(np.hstack([Xb,HOperm])):.3f}  (vs real {nho(np.hstack([Xb,HO])):.3f}; gain real if > null)")
print("\n  CANDIDATE 5 (attractor convergence): is the connectome more determined at adult than L1?")
# dispersion ratio per stage set (already: L1~0.44 adult~0.47) - recompute agreement fraction normalized
def stagerate(ds):
    ys=[np.array([1 if d in dchem.get((neurons[I[k]],neurons[J[k]]),set()) else 0 for k in range(N)]) for d in ds]
    k=np.sum(ys,0); n=len(ds); f=k/n
    var_ratio=(f*(1-f)).mean()/(f.mean()*(1-f.mean()))  # within-pair var / if-all-same-rate
    return f.mean(), var_ratio
er,vr=stagerate([1,2,3,4]); print(f"    L1(1-4): edge_rate={er:.3f} within/between var ratio={vr:.2f}")
er2,vr2=stagerate([7,8]); print(f"    adult(7,8): edge_rate={er2:.3f} within/between var ratio={vr2:.2f}")
print("    (lower ratio = more bimodal/determined; convergence would show adult << L1)")
