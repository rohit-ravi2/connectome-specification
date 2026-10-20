import openpyxl, csv, re, math, sys, numpy as np, pandas as pd
from collections import defaultdict
def P(*a): print(*a); sys.stdout.flush()
D="../data/celegans/data/connectome"
SC="../work"
# ---- neuron universe from adult nerve-ring contact xlsx (defines idx & Brittin area A) ----
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True))
neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T)
# ---- per-animal witvliet chemical edges (TARGET; W7/W8 adult replicates) ----
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
            if r['pre'] in idx and r['post'] in idx: dchem[(r['pre'],r['post'])].add(ds)
# ---- Cook adult chemical connectome C (READOUT source: connectome-from-connectome) ----
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r])); colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
P(f"Cook C edges={int(C.sum())}  M={M}")
Cin=C.T@C;Cout=C@C.T;Cff=C@C;Cu=((C+C.T)>0).astype(float);comm=Cu@Cu
Usv,s,Vt=np.linalg.svd(C); dd=16; RD=(Usv[:,:dd]*np.sqrt(s[:dd]))@(np.sqrt(s[:dd])[:,None]*Vt[:dd])
rho=np.abs(np.linalg.eigvals(C)).max(); beta=0.5/rho
Katz=np.linalg.inv(np.eye(M)-beta*C)-np.eye(M); C3=C@C@C
# ---- master (Witvliet per-stage contact + Sulston timing); UNIVERSE = master pairs ----
mst=pd.read_csv("../seed/data/processed/contact_decomposition/contact_decomposition_master.csv")
mst=mst[mst.na.isin(idx)&mst.nb.isin(idx)].reset_index(drop=True)
# adult witvliet contact adjacency (from c_adult) for contact-graph degree/shared-nbr
Cad=np.zeros((M,M))
for r in mst.itertuples():
    if not np.isnan(r.c_adult): Cad[idx[r.na],idx[r.nb]]=r.c_adult
Cad=np.maximum(Cad,Cad.T); Cadb=(Cad>0).astype(float); wdeg=Cadb.sum(1)
# node-level attrs
node_nt={};node_mod={};node_birth={}
for r in mst.itertuples():
    if isinstance(r.pre_nt,str): node_nt.setdefault(r.na,r.pre_nt)
    if isinstance(r.post_nt,str): node_nt.setdefault(r.nb,r.post_nt)
    if isinstance(r.pre_mod,str): node_mod.setdefault(r.na,r.pre_mod)
    if isinstance(r.post_mod,str): node_mod.setdefault(r.nb,r.post_mod)
    if not np.isnan(r.pre_birth): node_birth.setdefault(r.na,r.pre_birth)
    if not np.isnan(r.post_birth): node_birth.setdefault(r.nb,r.post_birth)
_stmap={'L1':1,'L2':2,'L3':3,'adult':4}
I=[];J=[];y7=[];y8=[];CONTACT=[];HO=[];TIMING=[]
for r in mst.itertuples():
    i,j=idx[r.na],idx[r.nb]; na,nb=r.na,r.nb
    st=dchem.get((na,nb),set())
    I.append(i);J.append(j);y7.append(1 if 7 in st else 0);y8.append(1 if 8 in st else 0)
    shared=int(((Cadb[i]>0)&(Cadb[j]>0)).sum())
    fcs=_stmap.get(r.first_contact_stage,np.nan)
    bdiff=abs(r.pre_birth-r.post_birth) if (not np.isnan(r.pre_birth) and not np.isnan(r.post_birth)) else np.nan
    # CONTACT (anatomy: Witvliet per-stage + Brittin area + contact-graph degree/shared) — non-circular
    CONTACT.append([r.c_L1,r.c_L2,r.c_L3,r.c_adult,math.log1p(max(r.c_adult,0) if not np.isnan(r.c_adult) else 0),
                    wdeg[i],wdeg[j],shared,math.log1p(A[i,j])])
    # HIGHER-ORDER (Cook adult connectome functions) => READOUT
    HO.append([np.log1p(Cin[i,j]),np.log1p(Cout[i,j]),np.log1p(Cff[i,j]),np.log1p(comm[i,j]),
               RD[i,j],RD[j,i],Katz[i,j],Katz[j,i],np.log1p(C3[i,j]),np.log1p(C3[j,i])])
    # DEVELOPMENTAL-TIMING (non-circular)
    TIMING.append([fcs,r.contact_persist,r.pre_birth,r.post_birth,r.both_born_preL1,bdiff])
I=np.array(I);J=np.array(J);y7=np.array(y7);y8=np.array(y8)
CONTACT=np.array(CONTACT,float);HO=np.array(HO,float);TIMING=np.array(TIMING,float)
def fill(X):
    X=X.copy()
    for c in range(X.shape[1]):
        col=X[:,c]; med=np.nanmedian(col); col[np.isnan(col)]=med; X[:,c]=col
    return X
missing_frac={'CONTACT':np.isnan(CONTACT).mean(),'HO':np.isnan(HO).mean(),'TIMING':np.isnan(TIMING).mean()}
CONTACT=fill(CONTACT);HO=fill(HO);TIMING=fill(TIMING)
CONTACT_NAMES=['c_L1','c_L2','c_L3','c_adult','log_c_adult','wcdeg_i','wcdeg_j','shared_contact_nbr','log_brittin_area']
HO_NAMES=['log_Cin(sharedIN)','log_Cout(sharedOUT)','log_Cff(feedfwd)','log_comm(commonNbr)','RDPG16_ij','RDPG16_ji','Katz_ij','Katz_ji','log_C3_ij','log_C3_ji']
TIMING_NAMES=['first_contact_stage','contact_persist','pre_birth','post_birth','both_born_preL1','abs_birth_diff']
yrep=((y7==1)&(y8==1)).astype(int)
P(f"pairs={len(I)}  yrep_rate={yrep.mean():.3f}  y7={y7.mean():.3f} y8={y8.mean():.3f} disagree={np.mean(y7!=y8):.3f}")
P("missing frac (pre-fill):",{k:round(float(v),3) for k,v in missing_frac.items()})
nt_arr=np.array([node_nt.get(n,'Unk') for n in neurons])
mod_arr=np.array([node_mod.get(n,'unk') for n in neurons])
birth_arr=np.array([node_birth.get(n,np.nan) for n in neurons],float)
np.savez(f"{SC}/feats.npz",I=I,J=J,y7=y7,y8=y8,yrep=yrep,
         CONTACT=CONTACT,HO=HO,TIMING=TIMING,
         CONTACT_NAMES=CONTACT_NAMES,HO_NAMES=HO_NAMES,TIMING_NAMES=TIMING_NAMES,
         neurons=np.array(neurons,dtype=object),C=C,A=A,cdeg=wdeg,
         nt_arr=nt_arr,mod_arr=mod_arr,birth_arr=birth_arr)
P("saved feats.npz")
