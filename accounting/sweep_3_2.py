import openpyxl, csv, re, math, json, numpy as np, glob
from collections import defaultdict
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.cluster import KMeans
from sklearn.model_selection import cross_val_predict, StratifiedKFold
from sklearn.metrics import roc_auc_score
D="../data/celegans/data/connectome"
wb=openpyxl.load_workbook(f"{D}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
data=list(wb['adult nerve ring neighbors'].iter_rows(values_only=True)); neurons=[h for h in data[0][1:] if isinstance(h,str)]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    if row[0] in idx:
        for j,v in enumerate(row[1:]):
            if v not in (None,''):
                try: A[idx[row[0]],j]=float(v)
                except: pass
A=np.maximum(A,A.T); cdeg=(A>0).sum(1)
dchem=defaultdict(set)
with open(f"{D}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        if r['type']=='chemical':
            ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1)); dchem[(r['pre'],r['post'])].add(ds)
def isneu(x): return bool(isinstance(x,str) and re.fullmatch(r"[A-Z]{2,4}[0-9]?[LR]?[0-9]?",str(x).strip()) and str(x).strip() in idx)
wc=openpyxl.load_workbook(f"{D}/cook2019/SI 5 Connectome adjacency matrices, corrected July 2020.xlsx",read_only=True,data_only=True)
cd=list(wc['hermaphrodite chemical'].iter_rows(values_only=True))
hr=max(range(6),key=lambda r:sum(isneu(x) for x in cd[r])); colmap={c:cd[hr][c].strip() for c in range(len(cd[hr])) if isneu(cd[hr][c])}
rhc=max(range(6),key=lambda c:sum(isneu(cd[r][c]) for r in range(hr+1,len(cd)) if c<len(cd[r])))
C=np.zeros((M,M))
for r in range(hr+1,len(cd)):
    if rhc<len(cd[r]) and isneu(cd[r][rhc]):
        pi=idx[cd[r][rhc].strip()]
        for c,post in colmap.items():
            if c<len(cd[r]) and cd[r][c] not in (None,''):
                try: C[pi,idx[post]]=1.0 if float(cd[r][c])>=1 else 0.0
                except: pass
Cin=C.T@C;Cout=C@C.T;Cff=C@C;Cu=((C+C.T)>0).astype(float);comm=Cu@Cu
W7=np.zeros((M,M))
for (a,b),s in dchem.items():
    if 7 in s and a in idx and b in idx: W7[idx[a],idx[b]]=1
W7in=W7.T@W7;W7out=W7@W7.T;W7ff=W7@W7;W7comm=((W7+W7.T)>0).astype(float)@((W7+W7.T)>0).astype(float)
# fascicle: skeleton centroids + voxel occupancy
SK="../seed/data/external/mechanistic_frontier/witvliet2021_developmental/data/skeletons"
sk=json.load(open(f"{SK}/Dataset7_skeletons.json")); cen={}; occ={}
for nm,v in sk.items():
    if nm not in idx: continue
    P=np.array(list(v['coords'].values()),float)
    if len(P)<3: continue
    cen[nm]=P.mean(0); occ[nm]=set(map(tuple,(P//2000).astype(int)))
names_sk=[n for n in neurons if n in cen]
cents=np.array([cen[n] for n in names_sk]); bund=KMeans(15,n_init=5,random_state=0).fit_predict(cents)
bundle={n:bund[i] for i,n in enumerate(names_sk)}
# activity: Atanas correlations
corsum=defaultdict(float); corcnt=defaultdict(int)
for f in glob.glob("../data/celegans/experiments/nb55_archetype_dynamics/data/upstream_artifacts/atanas_worm_*.npz"):
    d=np.load(f,allow_pickle=True); tr=d['neural']; ids=[str(x) for x in d['neuron_ids']]
    valid=[k for k,n in enumerate(ids) if n in idx]
    if len(valid)<5: continue
    Cc=np.corrcoef(tr[:,valid].T)
    for a in range(len(valid)):
        for b in range(a+1,len(valid)):
            key=(min(ids[valid[a]],ids[valid[b]]),max(ids[valid[a]],ids[valid[b]])); corsum[key]+=abs(Cc[a,b]); corcnt[key]+=1
def actcorr(ni,nj): 
    k=(min(ni,nj),max(ni,nj)); return corsum[k]/corcnt[k] if corcnt[k] else 0.0
print(f"fascicle neurons: {len(bundle)} | activity pairs covered: {len(corsum)}")
# pairs
I,J,y7,y8,feat,HO,FAS,ACT=[],[],[],[],[],[],[],[]
for i in range(M):
    for j in range(M):
        if i==j or A[i,j]<=0: continue
        s=dchem.get((neurons[i],neurons[j]),set()); ni,nj=neurons[i],neurons[j]
        I.append(i);J.append(j);y7.append(1 if 7 in s else 0);y8.append(1 if 8 in s else 0)
        feat.append([math.log1p(A[i,j]),int(((A[i]>0)&(A[j]>0)).sum()),float(A[i]@A[j]/(np.linalg.norm(A[i])*np.linalg.norm(A[j])+1e-9)),int(cdeg[i]),int(cdeg[j])])
        HO.append([np.log1p(Cin[i,j]),np.log1p(Cout[i,j]),np.log1p(Cff[i,j]),np.log1p(Cff[j,i]),np.log1p(comm[i,j]),np.log1p(W7in[i,j]),np.log1p(W7out[i,j]),np.log1p(W7ff[i,j]),np.log1p(W7comm[i,j])])
        sb=1.0 if (ni in bundle and nj in bundle and bundle[ni]==bundle[nj]) else 0.0
        cdst=float(np.linalg.norm(cen[ni]-cen[nj])) if ni in cen and nj in cen else 1e6
        jac=(len(occ[ni]&occ[nj])/max(1,len(occ[ni]|occ[nj]))) if ni in occ and nj in occ else 0.0
        FAS.append([sb,math.log1p(cdst),jac]); ACT.append([actcorr(ni,nj)])
I=np.array(I);J=np.array(J);y7=np.array(y7);y8=np.array(y8)
Xb=np.array(feat);HO=np.array(HO);FAS=np.array(FAS);ACT=np.array(ACT);N=len(I)
disagree=(y7!=y8).astype(int);yrep=((y7==1)&(y8==1)).astype(int);cv=StratifiedKFold(5,shuffle=True,random_state=0)
nf={n:i%5 for i,n in enumerate(sorted(set([neurons[i] for i in I]+[neurons[j] for j in J])))}
def nho(X):
    pred=np.full(N,np.nan)
    for f in range(5):
        te=[k for k in range(N) if nf[neurons[I[k]]]==f and nf[neurons[J[k]]]==f];tr=[k for k in range(N) if nf[neurons[I[k]]]!=f and nf[neurons[J[k]]]!=f]
        if len(te)<20: continue
        pred[np.array(te)]=GradientBoostingClassifier(random_state=0).fit(X[tr],yrep[tr]).predict_proba(X[te])[:,1]
    m=~np.isnan(pred);return roc_auc_score(yrep[m],pred[m])
def disp(X):
    y2=np.concatenate([y7,y8]);X2=np.vstack([X,X]);p=np.clip(cross_val_predict(GradientBoostingClassifier(random_state=0),X2,y2,cv=cv,method='predict_proba')[:,1][:N],1e-4,1-1e-4)
    return disagree.mean()/(2*p*(1-p)).mean()
BASE=np.hstack([Xb,HO])  # contact + higher-order (current best)
print("\nCandidates 3 (fascicle) + 2 (activity), ON TOP of contact+higher-order baseline")
print(f"  baseline (contact + higher-order):  disp={disp(BASE):.2f} AUC={nho(BASE):.3f}")
print(f"  + FASCICLE (bundle/position, cand3): disp={disp(np.hstack([BASE,FAS])):.2f} AUC={nho(np.hstack([BASE,FAS])):.3f}")
print(f"  + ACTIVITY (Atanas corr, cand2):     disp={disp(np.hstack([BASE,ACT])):.2f} AUC={nho(np.hstack([BASE,ACT])):.3f}")
print(f"  + BOTH:                              disp={disp(np.hstack([BASE,FAS,ACT])):.2f} AUC={nho(np.hstack([BASE,FAS,ACT])):.3f}")
print(f"  (activity is adult=circular caveat; fascicle may overlap contact)")
