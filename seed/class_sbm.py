"""FINAL TEST: is the reproducible connectome a neuron-CLASS address code, or sub-class PAIRWISE?
Assign each neuron its canonical class (Cook SI6, longest-prefix). Then:
 (1) class-level SBM: predict connection from class-pair (CV held-out AUC/residual).
 (2) within-class-pair consistency: do same-class-pair member neurons wire alike? (robust)."""
import openpyxl, csv, re, math, numpy as np
from collections import defaultdict
from sklearn.model_selection import StratifiedKFold
from sklearn.metrics import roc_auc_score

DIR="../data/celegans/data/connectome"
# ---- canonical class list from SI6 ----
wb=openpyxl.load_workbook(f"{DIR}/cook2019/SI 6 Cell class lists.xlsx",read_only=True,data_only=True)
ws=wb['cell class lists']; classes=set()
for row in ws.iter_rows(values_only=True):
    for c in row:
        if isinstance(c,str) and re.fullmatch(r'[A-Z]{2,4}[0-9]?',c.strip()):
            classes.add(c.strip())
classes=sorted(classes,key=len,reverse=True)  # longest first for prefix match
def to_class(n):
    if not isinstance(n,str): return None
    for c in classes:
        if n.startswith(c): return c
    return None

# ---- contact matrix + neurons ----
wb2=openpyxl.load_workbook(f"{DIR}/nerve_ring_neighbors/Adult and L4 nerve ring neighbors.xlsx",read_only=True,data_only=True)
ws2=wb2['adult nerve ring neighbors']; data=list(ws2.iter_rows(values_only=True))
neurons=[h for h in data[0][1:]]; idx={n:i for i,n in enumerate(neurons)}; M=len(neurons)
A=np.zeros((M,M))
for row in data[1:]:
    rn=row[0]
    if rn is None or rn not in idx: continue
    for j,val in enumerate(row[1:]):
        if val in (None,''): continue
        try: A[idx[rn],j]=float(val)
        except (TypeError,ValueError): pass
A=np.maximum(A,A.T)
ncls=[to_class(n) for n in neurons]
unmapped=[neurons[i] for i in range(M) if ncls[i] is None]
print(f"nerve-ring neurons: {M}   canonical classes in list: {len(classes)}")
print(f"mapped to a class: {M-len(unmapped)}/{M}   distinct classes present: {len(set(c for c in ncls if c))}")
if unmapped: print(f"  unmapped (kept as own class): {unmapped[:15]}")
ncls=[c if c else (neurons[i] if isinstance(neurons[i],str) else "NA") for i,c in enumerate(ncls)]  # fallback: own name

# ---- reproducible connectome ----
chem=defaultdict(set); elec=defaultdict(set)
with open(f"{DIR}/witvliet2020/witvliet_combined_edges.csv") as f:
    for r in csv.DictReader(f):
        ds=int(re.search(r'witvliet_2020_(\d)',r['source_file']).group(1))
        if ds not in (7,8): continue
        pre,post,typ=r['pre'],r['post'],r['type']
        (elec if typ=='electrical' else chem)[((min(pre,post),max(pre,post)) if typ=='electrical' else (pre,post))].add(ds)
repro=set()
for e,dss in chem.items():
    if len(dss)==2: repro.add((min(e),max(e)))
for e,dss in elec.items():
    if len(dss)==2: repro.add(e)

# contacting pairs with class-pair + label
cps=[]; y=[]
for i in range(M):
    for j in range(i+1,M):
        if A[i,j]<=0 or neurons[i] is None or neurons[j] is None: continue
        cp=tuple(sorted((ncls[i],ncls[j])))
        cps.append(cp); y.append(1 if (neurons[i],neurons[j]) in repro or (neurons[j],neurons[i]) in repro else 0)
y=np.array(y,float); N=len(y); pbar=y.mean()
base=N*(-(pbar*math.log2(pbar)+(1-pbar)*math.log2(1-pbar)))

# ---- (1) class-level SBM, CV ----
cv=StratifiedKFold(5,shuffle=True,random_state=0); preds=np.zeros(N); alpha=2.0
cps_arr=np.array(cps,dtype=object)
for tr,te in cv.split(np.zeros(N),y):
    tab=defaultdict(lambda:[0.0,0.0])
    for k in tr: tab[cps[k]][0]+=1; tab[cps[k]][1]+=y[k]
    for k in te:
        n,e=tab.get(cps[k],[0.0,0.0]); preds[k]=(e+alpha*pbar)/(n+alpha)
p=np.clip(preds,1e-6,1-1e-6)
sbm_resid=-(y*np.log2(p)+(1-y)*np.log2(1-p)).sum(); sbm_auc=roc_auc_score(y,p)

# ---- (2) within-class-pair consistency (multi-member class-pairs only) ----
groups=defaultdict(list)
for k in range(N): groups[cps[k]].append(y[k])
multi=[(g,v) for g,v in groups.items() if len(v)>=3]
within_bits=sum(len(v)*(-(np.mean(v)*math.log2(np.mean(v))+(1-np.mean(v))*math.log2(1-np.mean(v))) if 0<np.mean(v)<1 else 0) for g,v in multi)
n_multi_pairs=sum(len(v) for g,v in multi)
pure=sum(1 for g,v in multi if np.mean(v) in (0.0,1.0))

print(f"\ncontacting pairs N={N}, reproducible {int(y.sum())} ({100*pbar:.0f}%), distinct class-pairs {len(groups)}")
print(f"  base H(connectome)              = {base:.0f} bits  ({base/M:.2f}/neuron)")
print(f"  [refs] geometry 2907 (AUC .838), per-neuron identity AUC .653, degree AUC .654")
print(f"\n  (1) CLASS-LEVEL SBM (held-out): AUC = {sbm_auc:.3f}   residual = {sbm_resid:.0f} bits ({sbm_resid/M:.2f}/neuron)")
print(f"  (2) WITHIN-class-pair (>=3 members: {len(multi)} class-pairs, {n_multi_pairs} member-pairs):")
print(f"        pure (all-connected or all-not): {pure}/{len(multi)} class-pairs")
print(f"        within-class-pair residual entropy = {within_bits:.0f} bits over {n_multi_pairs} member-pairs "
      f"({within_bits/n_multi_pairs:.2f} bits/pair; base {(-(pbar*math.log2(pbar)+(1-pbar)*math.log2(1-pbar))):.2f})")
print(f"\n  VERDICT:")
if sbm_auc>0.80:
    print(f"    class-pair predicts wiring (AUC {sbm_auc:.2f}) -> neuron-CLASS is largely the address code.")
else:
    print(f"    class-pair only reaches AUC {sbm_auc:.2f}; within-class-pair members disagree "
          f"({within_bits/n_multi_pairs:.2f} vs base {(-(pbar*math.log2(pbar)+(1-pbar)*math.log2(1-pbar))):.2f} bits/pair)")
    print(f"    -> substantial SUB-CLASS PAIRWISE specificity beyond the neuron-class code.")
